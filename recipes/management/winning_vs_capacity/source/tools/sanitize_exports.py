import re,sys,html
from pptx import Presentation
SP,DL,OUT=sys.argv[1],sys.argv[2],sys.argv[3]
src=open(SP+'/export.py').read(); ns={}; exec(src[src.index('GEN='):src.index('def notes')],ns); GEN=ns['GEN']
NAME='Winning, and Increasing the Capacity to Win'
KIOSK=[('watch the kiosk being used','watch the product being used'),('support pre-stages kiosks','support pre-stages hardware')]
BAD=re.compile(r'mashgin|aramark|lafc|stanford|kiosk|micros|kapi|origin db|\bmpu\b',re.I)
def fix(t):
    if 'Mashgin examples:' in t:
        a=t.index('Mashgin examples:'); b=t.index('Caveat:',a)
        t=t[:a]+'Examples:\n'+'\n'.join('%d. %s'%(n+1,g) for n,g in enumerate(GEN))+'\n\n'+t[b:]
    for x,y in KIOSK: t=t.replace(x,y)
    return t
# PPTX
p=Presentation(f'{DL}/{NAME}.pptx')
for s in p.slides:
    if s.has_notes_slide:
        tf=s.notes_slide.notes_text_frame; old=tf.text; new=fix(old)
        if new!=old: tf.text=new
ids=p.slides._sldIdLst; last=ids[-1]
assert 'TODO' in ''.join(sh.text_frame.text for sh in p.slides[-1].shapes if sh.has_text_frame)
p.part.drop_rel(last.rId); ids.remove(last)
p.save(f'{OUT}/{NAME}.pptx')
for s in Presentation(f'{OUT}/{NAME}.pptx').slides:
    if s.has_notes_slide: assert not BAD.search(s.notes_slide.notes_text_frame.text)
# HTML
h=open(f'{DL}/{NAME}.html',encoding='utf-8').read()
h=re.sub(r'data-speaker-notes="([^"]*)"',lambda m:'data-speaker-notes="'+fix(m.group(1))+'"',h)
a=h.index('<section class="deck-slide" id="42"'); b=h.index('</section>',a)+len('</section>')
h=h[:a]+h[b:]
h=h.replace(' of 42"',' of 41"')
for m in re.finditer(r'data-speaker-notes="([^"]*)"',h): assert not BAD.search(m.group(1)),m.group(1)[:80]
assert 'TODO (not in the talk yet)' not in h
open(f'{OUT}/{NAME}.html','w',encoding='utf-8').write(h)
print('ok')
